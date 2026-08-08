"""Exportación de informes a CSV, Excel y PDF."""
from __future__ import annotations

import csv
import io
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from contaflow.config import APP_NAME, DIR_BACKUPS, RUTA_BD
from contaflow.services.utils import formato_moneda

AZUL = "1F3864"
GRIS = "F2F2F2"


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------


def a_csv(encabezados: list[str], filas: list[list], delimitador: str = ";") -> bytes:
    """CSV con BOM y punto y coma: se abre correctamente en Excel en español."""
    buffer = io.StringIO()
    escritor = csv.writer(buffer, delimiter=delimitador, lineterminator="\r\n")
    escritor.writerow(encabezados)
    for fila in filas:
        escritor.writerow(fila)
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def a_excel(
    titulo: str,
    encabezados: list[str],
    filas: list[list],
    *,
    subtitulo: str = "",
    columnas_monto: set[int] | None = None,
    hojas_extra: list[tuple[str, list[str], list[list]]] | None = None,
) -> bytes:
    """Genera un libro Excel con formato de informe contable."""
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31] or "Informe"

    borde = Border(*(Side(style="thin", color="D0D0D0"),) * 4)
    fila_actual = 1

    ws.cell(row=1, column=1, value=titulo).font = Font(size=14, bold=True, color=AZUL)
    fila_actual = 2
    if subtitulo:
        ws.cell(row=2, column=1, value=subtitulo).font = Font(size=10, italic=True, color="666666")
        fila_actual = 3
    fila_actual += 1

    encabezado_fila = fila_actual
    for col, texto in enumerate(encabezados, start=1):
        celda = ws.cell(row=encabezado_fila, column=col, value=texto)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde

    montos = columnas_monto or set()
    for i, fila in enumerate(filas, start=encabezado_fila + 1):
        for col, valor in enumerate(fila, start=1):
            celda = ws.cell(row=i, column=col, value=valor)
            celda.border = borde
            if col - 1 in montos and isinstance(valor, (int, float)):
                celda.number_format = "#,##0"
                celda.alignment = Alignment(horizontal="right")
        if i % 2 == 0:
            for col in range(1, len(encabezados) + 1):
                ws.cell(row=i, column=col).fill = PatternFill("solid", fgColor=GRIS)

    for col, texto in enumerate(encabezados, start=1):
        largo = max([len(str(texto))] + [
            len(str(f[col - 1])) for f in filas if len(f) >= col
        ][:200] or [10])
        ws.column_dimensions[get_column_letter(col)].width = min(45, max(11, largo + 3))
    ws.freeze_panes = ws.cell(row=encabezado_fila + 1, column=1)

    for nombre, cabeceras, datos in (hojas_extra or []):
        hoja = wb.create_sheet(nombre[:31])
        for col, texto in enumerate(cabeceras, start=1):
            celda = hoja.cell(row=1, column=col, value=texto)
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = PatternFill("solid", fgColor=AZUL)
        for i, fila in enumerate(datos, start=2):
            for col, valor in enumerate(fila, start=1):
                hoja.cell(row=i, column=col, value=valor)
        for col, texto in enumerate(cabeceras, start=1):
            hoja.column_dimensions[get_column_letter(col)].width = max(14, len(str(texto)) + 3)

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def a_pdf(
    titulo: str,
    encabezados: list[str],
    filas: list[list],
    *,
    subtitulo: str = "",
    empresa: str = "",
    apaisado: bool = False,
    alineacion_derecha: set[int] | None = None,
    notas: list[str] | None = None,
) -> bytes:
    """Informe en PDF listo para imprimir o adjuntar."""
    salida = io.BytesIO()
    tamano = landscape(A4) if apaisado else A4
    doc = SimpleDocTemplate(
        salida, pagesize=tamano,
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=14 * mm, bottomMargin=14 * mm,
        title=titulo, author=APP_NAME,
    )
    estilos = getSampleStyleSheet()
    est_titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=15, spaceAfter=2,
                                textColor=colors.HexColor("#1F3864"))
    est_sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=9,
                             textColor=colors.HexColor("#555555"), spaceAfter=8)
    est_celda = ParagraphStyle("c", parent=estilos["Normal"], fontSize=7.5, leading=9)
    est_cab = ParagraphStyle("h", parent=estilos["Normal"], fontSize=7.5, leading=9,
                             textColor=colors.white, fontName="Helvetica-Bold")

    elementos = [Paragraph(titulo, est_titulo)]
    if empresa:
        elementos.append(Paragraph(empresa, est_sub))
    if subtitulo:
        elementos.append(Paragraph(subtitulo, est_sub))
    elementos.append(Spacer(1, 4))

    datos = [[Paragraph(str(h), est_cab) for h in encabezados]]
    for fila in filas:
        datos.append([Paragraph(str(v if v is not None else ""), est_celda) for v in fila])

    tabla = Table(datos, repeatRows=1, hAlign="LEFT")
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
    ]
    for col in (alineacion_derecha or set()):
        estilo.append(("ALIGN", (col, 1), (col, -1), "RIGHT"))
    tabla.setStyle(TableStyle(estilo))
    elementos.append(tabla)

    if notas:
        elementos.append(Spacer(1, 10))
        for nota in notas:
            elementos.append(Paragraph(f"• {nota}", est_sub))

    pie = f"Generado por {APP_NAME} · {datetime.now():%d-%m-%Y %H:%M}"
    elementos.append(Spacer(1, 8))
    elementos.append(Paragraph(pie, est_sub))

    doc.build(elementos)
    return salida.getvalue()


def pdf_formulario(
    titulo: str, empresa: str, periodo: str, secciones: dict[str, list], *,
    resumen: list[tuple[str, str]] | None = None, notas: list[str] | None = None,
) -> bytes:
    """PDF de una propuesta de formulario, agrupado por sección y código."""
    salida = io.BytesIO()
    doc = SimpleDocTemplate(
        salida, pagesize=A4,
        leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
        title=titulo, author=APP_NAME,
    )
    estilos = getSampleStyleSheet()
    est_titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=16,
                                textColor=colors.HexColor("#1F3864"), spaceAfter=2)
    est_sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=9.5,
                             textColor=colors.HexColor("#444444"), spaceAfter=10)
    est_sec = ParagraphStyle("sec", parent=estilos["Heading3"], fontSize=10.5,
                             textColor=colors.HexColor("#1F3864"), spaceBefore=10, spaceAfter=4)
    est_nota = ParagraphStyle("n", parent=estilos["Normal"], fontSize=8,
                              textColor=colors.HexColor("#8A6D3B"), spaceAfter=3)

    elementos = [Paragraph(titulo, est_titulo), Paragraph(f"{empresa} · {periodo}", est_sub)]

    for seccion, codigos in secciones.items():
        elementos.append(Paragraph(seccion, est_sec))
        datos = [["Código", "Concepto", "Valor"]]
        for c in codigos:
            datos.append([c.codigo, c.etiqueta + (f"\n{c.nota}" if c.nota else ""),
                          formato_moneda(c.valor)])
        tabla = Table(datos, colWidths=[20 * mm, 118 * mm, 32 * mm], repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
            ("ALIGN", (2, 1), (2, -1), "RIGHT"),
            ("ALIGN", (0, 0), (0, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F7FA")]),
        ]))
        elementos.append(tabla)

    if resumen:
        elementos.append(Spacer(1, 12))
        datos = [[k, v] for k, v in resumen]
        tabla = Table(datos, colWidths=[110 * mm, 60 * mm])
        tabla.setStyle(TableStyle([
            ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 9.5),
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EAF0F8")),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AABBD4")),
            ("ALIGN", (1, 0), (1, -1), "RIGHT"),
        ]))
        elementos.append(tabla)

    if notas:
        elementos.append(Spacer(1, 12))
        for nota in notas:
            elementos.append(Paragraph(f"⚠ {nota}", est_nota))

    elementos.append(Spacer(1, 10))
    elementos.append(Paragraph(
        f"Generado por {APP_NAME} · {datetime.now():%d-%m-%Y %H:%M} · "
        "Documento de trabajo, no constituye declaración ante el SII.",
        ParagraphStyle("f", parent=estilos["Normal"], fontSize=7.5,
                       textColor=colors.HexColor("#888888")),
    ))
    doc.build(elementos)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# Respaldos
# ---------------------------------------------------------------------------


def crear_respaldo() -> Path:
    """Copia la base de datos a la carpeta de respaldos."""
    destino = DIR_BACKUPS / f"contaflow-{datetime.now():%Y%m%d-%H%M%S}.db"
    shutil.copy2(RUTA_BD, destino)
    return destino


def listar_respaldos() -> list[Path]:
    return sorted(DIR_BACKUPS.glob("contaflow-*.db"), reverse=True)


def restaurar_respaldo(ruta: Path) -> None:
    """Restaura un respaldo (guardando antes el estado actual)."""
    if not ruta.exists():
        raise FileNotFoundError(f"No existe el respaldo {ruta}")
    previo = DIR_BACKUPS / f"antes-de-restaurar-{datetime.now():%Y%m%d-%H%M%S}.db"
    shutil.copy2(RUTA_BD, previo)
    shutil.copy2(ruta, RUTA_BD)
